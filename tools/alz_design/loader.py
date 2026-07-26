from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .common import Context, DesignError, index_by, is_enabled, key_values, resolve_value, rows

REQUIRED_SHEETS = [
    "01_Global", "02_ResourceGroups", "03_HubNetwork", "04_HubSubnets", "05_Workloads",
    "06_WorkloadSubnets", "07_VMs", "08_VMDisks", "09_AKSClusters", "10_AKSNodePools",
    "11_AIServices", "12_PrivateDNSZones", "99_GenerationMap",
]


def load_context(workbook_path: Path, allow_placeholders: bool) -> tuple[Context, list[dict[str, Any]]]:
    wb = load_workbook(workbook_path, data_only=True)
    missing = [name for name in REQUIRED_SHEETS if name not in wb.sheetnames]
    if missing:
        raise DesignError(f"missing sheets: {', '.join(missing)}")

    globals_raw = key_values(wb["01_Global"])
    globals_resolved = {key: resolve_value(value, allow_placeholders) for key, value in globals_raw.items()}
    ctx = Context(
        global_values=globals_resolved,
        resource_groups=index_by(rows(wb["02_ResourceGroups"]), "resource_group_id", "02_ResourceGroups"),
        hubs=index_by(rows(wb["03_HubNetwork"]), "hub_id", "03_HubNetwork"),
        hub_subnets=rows(wb["04_HubSubnets"]),
        workloads=index_by(rows(wb["05_Workloads"]), "workload_id", "05_Workloads"),
        workload_subnets=index_by(rows(wb["06_WorkloadSubnets"]), "subnet_id", "06_WorkloadSubnets"),
        vms=index_by(rows(wb["07_VMs"]), "vm_id", "07_VMs"),
        vm_disks=rows(wb["08_VMDisks"]),
        aks_clusters=index_by(rows(wb["09_AKSClusters"]), "cluster_id", "09_AKSClusters"),
        aks_node_pools=rows(wb["10_AKSNodePools"]),
        ai_services=index_by(rows(wb["11_AIServices"]), "ai_id", "11_AIServices"),
        private_dns_rows=rows(wb["12_PrivateDNSZones"]),
    )
    generation_map = [row for row in rows(wb["99_GenerationMap"]) if is_enabled(row)]
    return ctx, generation_map
