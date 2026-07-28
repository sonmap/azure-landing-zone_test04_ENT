#!/usr/bin/env python3
"""Create the source-of-truth Excel workbook used to generate Terraform auto tfvars."""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
REQUIRED_FILL = PatternFill("solid", fgColor="FFF2CC")


def add_sheet(wb: Workbook, name: str, headers: list[str], rows: list[list[object]]) -> None:
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
        ws.column_dimensions[column[0].column_letter].width = width


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", default="design/azure_landingzone_design.xlsx")
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()

    output = Path(args.out)
    if output.exists() and not args.force:
        raise SystemExit(f"File already exists: {output}. Use --force to replace it.")
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    add_sheet(
        wb,
        "00_Control",
        ["key", "value", "description"],
        [
            ["design_version", "1.1", "Excel design schema version"],
            ["design_status", "DRAFT", "DRAFT/APPROVED"],
            ["owner", "cloud-platform", "Design owner"],
        ],
    )
    add_sheet(
        wb,
        "01_Global",
        ["key", "value", "description"],
        [
            ["tenant_id", "${TENANT_ID}", "Resolved from --tenant-id or TENANT_ID"],
            ["subscription_id", "${SUBSCRIPTION_ID}", "Resolved from --subscription-id or SUBSCRIPTION_ID"],
            ["default_location", "koreacentral", "Default Azure region"],
            ["project", "samsunglife-landingzone", "Common project tag"],
            ["managed_by", "terraform", "Common management tag"],
            ["hub_resource_group_name", "rg-sl-hub-krc", "Hub resource group"],
            ["hub_vnet_name", "vnet-sl-hub-krc", "Hub VNet"],
            ["hub_dns_inbound_ip", "10.39.1.4", "DNS resolver inbound IP"],
            ["firewall_private_ip", "10.39.0.68", "Azure Firewall private IP"],
            ["ssh_public_key", "${SSH_PUBLIC_KEY}", "Resolved from --ssh-public-key or SSH_PUBLIC_KEY"],
        ],
    )
    add_sheet(
        wb,
        "02_ResourceGroups",
        ["resource_group_id", "enabled", "name", "location", "environment", "department", "owner", "costcenter"],
        [
            ["rg-hub", "Y", "rg-sl-hub-krc", "koreacentral", "shared", "platform", "cloud-team", "CC10000"],
            ["rg-sales-dev", "Y", "rg-sl-sales-dev", "koreacentral", "dev", "sales", "sales-it", "CC10020"],
        ],
    )
    add_sheet(
        wb,
        "03_HubNetwork",
        ["hub_id", "enabled", "resource_group_id", "vnet_name", "location", "address_space", "dns_servers", "private_dns_resolver"],
        [["hub-krc", "Y", "rg-hub", "vnet-sl-hub-krc", "koreacentral", "10.39.0.0/20", "", "Y"]],
    )
    add_sheet(
        wb,
        "04_HubSubnets",
        ["subnet_id", "enabled", "hub_id", "name", "address_prefix", "create_nsg", "delegate_dns_resolver", "dns_endpoint_role"],
        [
            ["gateway", "Y", "hub-krc", "GatewaySubnet", "10.39.0.0/27", "N", "N", ""],
            ["firewall", "Y", "hub-krc", "AzureFirewallSubnet", "10.39.0.64/26", "N", "N", ""],
            ["dns_in", "Y", "hub-krc", "snet-dns-inbound", "10.39.1.0/28", "N", "Y", "inbound"],
            ["dns_out", "Y", "hub-krc", "snet-dns-outbound", "10.39.1.16/28", "N", "Y", "outbound"],
            ["bastion", "Y", "hub-krc", "AzureBastionSubnet", "10.39.2.0/26", "N", "N", ""],
            ["pe", "Y", "hub-krc", "snet-shared-pe", "10.39.3.0/24", "Y", "N", ""],
        ],
    )
    add_sheet(
        wb,
        "05_Workloads",
        ["workload_id", "enabled", "department", "workload", "environment", "resource_group_id", "vnet_name", "address_space", "location"],
        [["sales-dev", "Y", "sales", "sales-support", "dev", "rg-sales-dev", "vnet-sl-sales-dev", "10.40.0.0/20", "koreacentral"]],
    )
    add_sheet(
        wb,
        "06_WorkloadSubnets",
        ["subnet_id", "enabled", "workload_id", "name", "address_prefix", "create_nsg", "associate_route_table", "private_endpoint_network_policies"],
        [
            ["sales-dev-web", "Y", "sales-dev", "snet-web", "10.40.1.0/24", "Y", "Y", "Enabled"],
            ["sales-dev-was", "Y", "sales-dev", "snet-was", "10.40.2.0/24", "Y", "Y", "Enabled"],
            ["sales-dev-db", "Y", "sales-dev", "snet-db", "10.40.3.0/24", "Y", "Y", "Enabled"],
            ["sales-dev-pe", "Y", "sales-dev", "snet-pe", "10.40.4.0/24", "Y", "N", "Disabled"],
            ["sales-dev-aks", "Y", "sales-dev", "snet-aks", "10.40.5.0/24", "Y", "Y", "Enabled"],
        ],
    )
    add_sheet(
        wb,
        "07_VMs",
        ["vm_id", "enabled", "workload_id", "role", "vm_name", "subnet_id", "private_ip", "vm_size", "os", "os_disk_gb", "admin_username", "itsm_ticket"],
        [
            ["sales-dev-web01", "Y", "sales-dev", "web", "vm-sl-sales-dev-web01", "sales-dev-web", "10.40.1.10", "Standard_D2s_v3", "RHEL9", 128, "azureuser", "CSR-VM-001"],
            ["sales-dev-was01", "Y", "sales-dev", "was", "vm-sl-sales-dev-was01", "sales-dev-was", "10.40.2.10", "Standard_D4s_v3", "RHEL9", 128, "azureuser", "CSR-VM-001"],
        ],
    )
    add_sheet(
        wb,
        "08_VMDisks",
        ["disk_id", "enabled", "vm_id", "disk_name", "size_gb", "lun", "storage_type", "caching"],
        [
            ["web01-app", "Y", "sales-dev-web01", "app", 128, 0, "Premium_LRS", "ReadWrite"],
            ["web01-log", "Y", "sales-dev-web01", "log", 256, 1, "Premium_LRS", "None"],
            ["was01-app", "Y", "sales-dev-was01", "app", 256, 0, "Premium_LRS", "ReadWrite"],
        ],
    )
    add_sheet(
        wb,
        "09_PrivateDNSZones",
        ["dns_zone_id", "enabled", "resource_group_id", "zone_name", "environment", "department", "owner", "itsm_ticket"],
        [
            ["pdns-internal", "Y", "rg-hub", "internal.sl.local", "shared", "platform", "cloud-team", "CSR-DNS-001"],
            ["pdns-blob", "N", "rg-hub", "privatelink.blob.core.windows.net", "shared", "platform", "cloud-team", "CSR-DNS-002"],
        ],
    )
    add_sheet(
        wb,
        "10_PrivateDNSLinks",
        ["link_id", "enabled", "dns_zone_id", "vnet_scope", "vnet_key", "link_name", "registration_enabled"],
        [
            ["link-internal-hub", "Y", "pdns-internal", "hub", "hub-krc", "link-internal-hub-krc", "N"],
            ["link-internal-sales", "Y", "pdns-internal", "workload", "sales-dev", "link-internal-sales-dev", "N"],
            ["link-blob-hub", "N", "pdns-blob", "hub", "hub-krc", "link-blob-hub-krc", "N"],
        ],
    )
    add_sheet(
        wb,
        "11_PrivateEndpoints",
        [
            "private_endpoint_id", "enabled", "resource_group_id", "workload_id", "subnet_id",
            "pe_name", "private_connection_resource_id", "subresource_names",
            "private_dns_zone_ids", "environment", "department", "itsm_ticket",
        ],
        [
            [
                "pe-sales-storage-blob", "N", "rg-sales-dev", "sales-dev", "sales-dev-pe",
                "pe-stslsalesdev001-blob",
                "/subscriptions/${SUBSCRIPTION_ID}/resourceGroups/rg-sl-sales-dev/providers/Microsoft.Storage/storageAccounts/stslsalesdev001",
                "blob", "pdns-blob", "dev", "sales", "CSR-PE-001",
            ],
        ],
    )
    add_sheet(
        wb,
        "12_FirewallRuleGroups",
        [
            "rule_group_id", "enabled", "firewall_policy_id", "rule_collection_group_name",
            "priority", "environment", "department", "itsm_ticket",
        ],
        [
            [
                "fw-sales-dev", "N",
                "/subscriptions/${SUBSCRIPTION_ID}/resourceGroups/rg-sl-hub-krc/providers/Microsoft.Network/firewallPolicies/fwpol-sl-hub-krc",
                "rcg-sales-dev", 120, "dev", "sales", "CSR-FW-001",
            ],
        ],
    )
    add_sheet(
        wb,
        "13_FirewallNetworkRules",
        [
            "rule_id", "enabled", "rule_group_id", "collection_id", "collection_name",
            "collection_priority", "action", "name", "protocols", "source_addresses",
            "destination_addresses", "destination_ports",
        ],
        [
            [
                "fw-net-oracle", "N", "fw-sales-dev", "app-to-db", "nrc-sales-was-to-db",
                100, "Allow", "sales-was-to-onprem-oracle", "TCP", "10.40.2.0/24",
                "10.10.20.30", "1521",
            ],
        ],
    )
    add_sheet(
        wb,
        "14_FirewallApplicationRules",
        [
            "rule_id", "enabled", "rule_group_id", "collection_id", "collection_name",
            "collection_priority", "action", "name", "source_addresses",
            "protocol_type", "protocol_port", "destination_fqdns",
        ],
        [
            [
                "fw-app-azure", "N", "fw-sales-dev", "allow-azure", "arc-sales-azure",
                200, "Allow", "sales-to-azure", "10.40.0.0/20", "Https", 443,
                "management.azure.com|login.microsoftonline.com",
            ],
        ],
    )
    add_sheet(
        wb,
        "15_DNSRecords",
        ["record_id", "enabled", "dns_zone_id", "name", "ttl", "records", "environment", "department", "itsm_ticket"],
        [
            ["sales-web", "Y", "pdns-internal", "sales-dev-web", 300, "10.40.1.10", "dev", "sales", "CSR-DNS-003"],
            ["sales-was", "Y", "pdns-internal", "sales-dev-was", 300, "10.40.2.10", "dev", "sales", "CSR-DNS-003"],
        ],
    )
    add_sheet(
        wb,
        "99_GenerationMap",
        ["generation_id", "enabled", "module_type", "source_key", "output_root", "output_filename"],
        [
            ["gen-foundation-rg", "Y", "resource_groups", "all", "00-foundation/resource-groups", "10-design.auto.tfvars.json"],
            ["gen-hub-krc", "Y", "hub_network", "hub-krc", "10-platform/hub-network", "10-design.auto.tfvars.json"],
            ["gen-sales-dev-spoke", "Y", "workload_spoke", "sales-dev", "20-workload/sales-dev-spoke", "10-design.auto.tfvars.json"],
            ["gen-sales-dev-vm", "Y", "vm", "sales-dev", "30-services/vm-sales-dev", "10-design.auto.tfvars.json"],
            ["gen-private-dns", "Y", "private_dns_zones", "rg-hub", "40-access/private-dns-zones", "10-design.auto.tfvars.json"],
            ["gen-firewall-rule", "N", "firewall_rule", "fw-sales-dev", "40-access/firewall-rule", "10-design.auto.tfvars.json"],
            ["gen-private-endpoint", "N", "private_endpoint", "pe-sales-storage-blob", "40-access/private-endpoint", "10-design.auto.tfvars.json"],
            ["gen-dns-record", "Y", "dns_records", "pdns-internal", "40-access/dns-record", "10-design.auto.tfvars.json"],
        ],
    )

    yes_no = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    for ws in wb.worksheets:
        headers = [c.value for c in ws[1]]
        if "enabled" in headers:
            col = headers.index("enabled") + 1
            yes_no.add(f"{ws.cell(2, col).coordinate}:{ws.cell(5000, col).coordinate}")
            ws.add_data_validation(yes_no)

    scope_validation = DataValidation(type="list", formula1='"hub,workload"', allow_blank=False)
    links_ws = wb["10_PrivateDNSLinks"]
    scope_validation.add("D2:D5000")
    links_ws.add_data_validation(scope_validation)

    wb.save(output)
    print(f"created: {output}")


if __name__ == "__main__":
    main()
