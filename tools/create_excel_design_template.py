#!/usr/bin/env python3
"""Create the Azure Landing Zone Excel design workbook template."""
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def add_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col_idx, header in enumerate(headers, 1):
        max_len = max(len(str(header)), *(len(str(ws.cell(r, col_idx).value or "")) for r in range(2, ws.max_row + 1)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 45)


def add_yes_no_validation(ws, column_letter: str, start: int = 2, end: int = 1000) -> None:
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}{start}:{column_letter}{end}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="design/azure_landingzone_design.xlsx")
    args = parser.parse_args()

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    add_sheet(
        wb,
        "00_Control",
        ["key", "value", "description"],
        [
            ["design_version", "1.0", "Excel 설계 버전"],
            ["design_name", "azure-landing-zone", "설계 이름"],
            ["approval_status", "DRAFT", "DRAFT/APPROVED"],
            ["output_filename", "10-design.auto.tfvars.json", "Terraform 자동 변수 파일명"],
        ],
    )

    add_sheet(
        wb,
        "01_Global",
        ["key", "value", "description"],
        [
            ["tenant_id", "${ENV:TENANT_ID}", "Azure Tenant ID"],
            ["subscription_id", "${ENV:SUBSCRIPTION_ID}", "Azure Subscription ID"],
            ["default_location", "koreacentral", "기본 Azure Region"],
            ["project", "samsunglife-landingzone", "공통 project tag"],
            ["managed_by", "terraform", "공통 managed_by tag"],
            ["hub_resource_group_name", "rg-sl-hub-krc", "Hub Resource Group"],
            ["hub_vnet_name", "vnet-sl-hub-krc", "Hub VNet"],
            ["hub_dns_inbound_ip", "10.39.1.4", "Private DNS Resolver inbound IP"],
            ["firewall_private_ip", "10.39.0.68", "Azure Firewall private IP"],
            ["ssh_public_key", "${ENV:SSH_PUBLIC_KEY}", "Linux SSH public key"],
            ["aks_log_analytics_workspace_id", "${ENV:AKS_LOG_ANALYTICS_WORKSPACE_ID}", "AKS Log Analytics workspace resource ID"],
        ],
    )

    add_sheet(
        wb,
        "02_ResourceGroups",
        ["resource_group_id", "enabled", "name", "location", "environment", "department", "owner", "costcenter", "data_class", "itsm_ticket"],
        [
            ["rg-hub", "Y", "rg-sl-hub-krc", "koreacentral", "shared", "platform", "cloud-team", "CC10000", "internal", "CSR-LZ-001"],
            ["rg-sales-dev", "Y", "rg-sl-sales-dev", "koreacentral", "dev", "sales", "sales-it", "CC10020", "internal", "CSR-SALES-001"],
        ],
    )

    add_sheet(
        wb,
        "03_HubNetwork",
        ["hub_id", "enabled", "resource_group_id", "vnet_name", "location", "address_space", "dns_servers", "enable_private_dns_resolver", "dns_inbound_subnet_id", "dns_outbound_subnet_id"],
        [["hub-krc", "Y", "rg-hub", "vnet-sl-hub-krc", "koreacentral", "10.39.0.0/20", "", "Y", "hub-dns-in", "hub-dns-out"]],
    )

    add_sheet(
        wb,
        "04_HubSubnets",
        ["subnet_id", "enabled", "hub_id", "name", "address_prefix", "create_nsg", "delegate_dns_resolver"],
        [
            ["hub-gateway", "Y", "hub-krc", "GatewaySubnet", "10.39.0.0/27", "N", "N"],
            ["hub-firewall", "Y", "hub-krc", "AzureFirewallSubnet", "10.39.0.64/26", "N", "N"],
            ["hub-dns-in", "Y", "hub-krc", "snet-dns-inbound", "10.39.1.0/28", "N", "Y"],
            ["hub-dns-out", "Y", "hub-krc", "snet-dns-outbound", "10.39.1.16/28", "N", "Y"],
            ["hub-bastion", "Y", "hub-krc", "AzureBastionSubnet", "10.39.2.0/26", "N", "N"],
            ["hub-pe", "Y", "hub-krc", "snet-shared-pe", "10.39.3.0/24", "Y", "N"],
        ],
    )

    add_sheet(
        wb,
        "05_Workloads",
        ["workload_id", "enabled", "department", "workload", "environment", "resource_group_id", "vnet_name", "location", "address_space", "dns_servers", "hub_id"],
        [["sales-dev", "Y", "sales", "sales-support", "dev", "rg-sales-dev", "vnet-sl-sales-dev", "koreacentral", "10.40.0.0/20", "10.39.1.4", "hub-krc"]],
    )

    add_sheet(
        wb,
        "06_WorkloadSubnets",
        ["subnet_id", "enabled", "workload_id", "name", "address_prefix", "create_nsg", "associate_route_table", "private_endpoint_network_policies"],
        [
            ["sales-dev-web", "Y", "sales-dev", "snet-web", "10.40.1.0/24", "Y", "Y", "Enabled"],
            ["sales-dev-was", "Y", "sales-dev", "snet-was", "10.40.2.0/24", "Y", "Y", "Enabled"],
            ["sales-dev-db", "Y", "sales-dev", "snet-db", "10.40.3.0/24", "Y", "Y", "Enabled"],
            ["sales-dev-pe", "Y", "sales-dev", "snet-pe", "10.40.4.0/24", "Y", "N", "Disabled"],
            ["sales-dev-aks", "Y", "sales-dev", "snet-aks", "10.40.5.0/24", "Y", "Y", "Enabled"],
        ],
    )

    add_sheet(
        wb,
        "07_VMs",
        ["vm_id", "enabled", "workload_id", "role", "vm_name", "subnet_id", "private_ip", "vm_size", "os", "os_disk_gb", "admin_username", "itsm_ticket"],
        [
            ["sales-dev-web01", "Y", "sales-dev", "web", "vm-sl-sales-dev-web01", "sales-dev-web", "10.40.1.10", "Standard_D2s_v5", "RHEL9", 128, "azureuser", "CSR-VM-001"],
            ["sales-dev-was01", "Y", "sales-dev", "was", "vm-sl-sales-dev-was01", "sales-dev-was", "10.40.2.10", "Standard_D4s_v5", "RHEL9", 128, "azureuser", "CSR-VM-001"],
        ],
    )

    add_sheet(
        wb,
        "08_VMDisks",
        ["disk_id", "enabled", "vm_id", "disk_name", "size_gb", "lun", "storage_type", "caching"],
        [
            ["web01-app", "Y", "sales-dev-web01", "app", 128, 0, "Premium_LRS", "ReadWrite"],
            ["web01-log", "Y", "sales-dev-web01", "log", 256, 1, "Premium_LRS", "None"],
            ["was01-app", "Y", "sales-dev-was01", "app", 256, 0, "Premium_LRS", "ReadWrite"],
        ],
    )

    add_sheet(
        wb,
        "09_AKSClusters",
        ["cluster_id", "enabled", "workload_id", "name", "subnet_id", "kubernetes_version", "dns_prefix", "private_cluster_enabled", "sku_tier", "local_account_disabled", "service_cidr", "dns_service_ip", "outbound_type", "network_plugin", "network_policy", "monitoring_enabled"],
        [["sales-dev-aks", "Y", "sales-dev", "aks-sl-sales-dev", "sales-dev-aks", "1.30", "aks-sl-sales-dev", "Y", "Free", "Y", "10.250.0.0/16", "10.250.0.10", "userDefinedRouting", "azure", "azure", "Y"]],
    )

    add_sheet(
        wb,
        "10_AKSNodePools",
        ["node_pool_id", "enabled", "cluster_id", "pool_type", "name", "vm_size", "mode", "node_count", "auto_scaling_enabled", "min_count", "max_count", "os_disk_size_gb"],
        [
            ["sales-dev-system", "Y", "sales-dev-aks", "default", "system", "Standard_D2s_v5", "System", 1, "Y", 1, 2, 64],
            ["sales-dev-app", "N", "sales-dev-aks", "user", "app", "Standard_D4s_v5", "User", 1, "Y", 1, 3, 128],
        ],
    )

    add_sheet(
        wb,
        "11_AIServices",
        ["ai_id", "enabled", "workload_id", "department", "environment", "private_endpoint_subnet_id", "private_dns_zone_resource_group_id", "openai_enabled", "openai_name", "openai_sku", "search_enabled", "search_name", "search_sku", "storage_enabled", "storage_name", "storage_replication", "keyvault_enabled", "keyvault_name", "purge_protection"],
        [["sales-ai-sandbox", "N", "sales-dev", "sales", "sandbox", "sales-dev-pe", "rg-hub", "Y", "oai-sl-sales-sbox-001", "S0", "Y", "srch-sl-sales-sbox-001", "basic", "Y", "stslssalesaisbox001", "LRS", "Y", "kv-sl-sales-sbox-001", "Y"]],
    )

    add_sheet(
        wb,
        "12_PrivateDNSZones",
        ["dns_set_id", "enabled", "resource_group_id", "zone_name", "link_id", "link_name", "hub_id", "registration_enabled"],
        [
            ["private-dns", "Y", "rg-hub", "privatelink.openai.azure.com", "hub-openai", "lnk-hub-openai", "hub-krc", "N"],
            ["private-dns", "Y", "rg-hub", "privatelink.search.windows.net", "hub-search", "lnk-hub-search", "hub-krc", "N"],
            ["private-dns", "Y", "rg-hub", "privatelink.blob.core.windows.net", "hub-blob", "lnk-hub-blob", "hub-krc", "N"],
            ["private-dns", "Y", "rg-hub", "privatelink.vaultcore.azure.net", "hub-keyvault", "lnk-hub-keyvault", "hub-krc", "N"],
        ],
    )

    add_sheet(
        wb,
        "99_GenerationMap",
        ["generation_id", "enabled", "module_type", "source_key", "output_root", "output_filename"],
        [
            ["gen-foundation-rg", "Y", "resource_groups", "all", "00-foundation/resource-groups", "10-design.auto.tfvars.json"],
            ["gen-hub-krc", "Y", "hub_network", "hub-krc", "10-platform/hub-network", "10-design.auto.tfvars.json"],
            ["gen-sales-dev-spoke", "Y", "workload_spoke", "sales-dev", "20-workload/sales-dev-spoke", "10-design.auto.tfvars.json"],
            ["gen-sales-dev-vm", "Y", "vm", "sales-dev", "30-services/vm-sales-dev", "10-design.auto.tfvars.json"],
            ["gen-sales-dev-aks", "Y", "aks", "sales-dev-aks", "30-services/aks-sales-dev", "10-design.auto.tfvars.json"],
            ["gen-sales-ai", "N", "ai", "sales-ai-sandbox", "30-services/ai-sales-sandbox", "10-design.auto.tfvars.json"],
            ["gen-private-dns", "Y", "private_dns_zones", "private-dns", "40-access/private-dns-zones", "10-design.auto.tfvars.json"],
        ],
    )

    for sheet in ["02_ResourceGroups", "03_HubNetwork", "04_HubSubnets", "05_Workloads", "06_WorkloadSubnets", "07_VMs", "08_VMDisks", "09_AKSClusters", "10_AKSNodePools", "11_AIServices", "12_PrivateDNSZones", "99_GenerationMap"]:
        ws = wb[sheet]
        enabled_col = next((cell.column_letter for cell in ws[1] if cell.value == "enabled"), None)
        if enabled_col:
            add_yes_no_validation(ws, enabled_col)

    wb.save(output)
    print(f"created: {output}")


if __name__ == "__main__":
    main()
