#!/usr/bin/env python3
"""Generate root-specific 10-design.auto.tfvars.json files from the Excel design workbook."""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
from collections import defaultdict
from ipaddress import ip_address, ip_network
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

YES = {"y", "yes", "true", "1", "enabled", "enable", "사용", "예"}


class DesignError(Exception):
    pass


def enabled(value: Any) -> bool:
    return str(value or "").strip().lower() in YES


def rows(ws) -> list[dict[str, Any]]:
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(v).strip() if v is not None else "" for v in values[0]]
    result: list[dict[str, Any]] = []
    for raw in values[1:]:
        if not any(v not in (None, "") for v in raw):
            continue
        result.append({headers[i]: raw[i] for i in range(min(len(headers), len(raw))) if headers[i]})
    return result


def kv_sheet(ws) -> dict[str, str]:
    result: dict[str, str] = {}
    for row in rows(ws):
        key = str(row.get("key", "")).strip()
        if key:
            result[key] = str(row.get("value", "")).strip()
    return result


def split_list(value: Any) -> list[str]:
    text = str(value or "").strip()
    return [part.strip() for part in re.split(r"[|,]", text) if part.strip()]


def require(row: dict[str, Any], *keys: str) -> None:
    missing = [key for key in keys if row.get(key) in (None, "")]
    if missing:
        raise DesignError(f"Missing required columns {missing}: {row}")


def unique(items: list[dict[str, Any]], key: str, sheet: str) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in items:
        if not enabled(row.get("enabled", "Y")):
            continue
        value = str(row.get(key, "")).strip()
        if not value:
            raise DesignError(f"{sheet}: empty {key}")
        if value in result:
            raise DesignError(f"{sheet}: duplicate {key}={value}")
        result[value] = row
    return result


def resolve_placeholders(value: Any, context: dict[str, str]) -> Any:
    if isinstance(value, str):
        def repl(match: re.Match[str]) -> str:
            key = match.group(1)
            return context.get(key, match.group(0))
        return re.sub(r"\$\{([A-Za-z0-9_]+)\}", repl, value)
    if isinstance(value, list):
        return [resolve_placeholders(v, context) for v in value]
    if isinstance(value, dict):
        return {k: resolve_placeholders(v, context) for k, v in value.items()}
    return value


def ensure_cidr(value: str, label: str) -> None:
    try:
        ip_network(value, strict=False)
    except ValueError as exc:
        raise DesignError(f"Invalid CIDR for {label}: {value}") from exc


def ensure_ip_in_subnet(ip_value: str, subnet_value: str, label: str) -> None:
    try:
        if ip_address(ip_value) not in ip_network(subnet_value, strict=False):
            raise DesignError(f"{label}: IP {ip_value} is not in subnet {subnet_value}")
    except ValueError as exc:
        raise DesignError(f"{label}: invalid IP/CIDR {ip_value}, {subnet_value}") from exc


def image_for(os_name: str) -> tuple[str, str, str]:
    text = os_name.lower()
    if "win" in text:
        return "MicrosoftWindowsServer", "WindowsServer", "2022-datacenter-g2"
    if "ubuntu" in text:
        return "Canonical", "0001-com-ubuntu-server-jammy", "22_04-lts-gen2"
    return "RedHat", "RHEL", "9-lvm-gen2"


def write_json(path: Path, data: dict[str, Any], context: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rendered = resolve_placeholders(data, context)
    unresolved = re.findall(r"\$\{[A-Za-z0-9_]+\}", json.dumps(rendered))
    if unresolved:
        raise DesignError(f"Unresolved placeholders in {path}: {sorted(set(unresolved))}")
    path.write_text(json.dumps(rendered, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"created: {path}")


def common_tags(global_cfg: dict[str, str], row: dict[str, Any] | None = None) -> dict[str, str]:
    tags = {
        "project": global_cfg["project"],
        "managed_by": global_cfg["managed_by"],
    }
    if row:
        for key in ("environment", "department", "owner", "costcenter", "itsm_ticket"):
            value = str(row.get(key, "") or "").strip()
            if value:
                tags[key] = value
    return tags


def resource_id(subscription_id: str, resource_group_name: str, provider_path: str) -> str:
    return (
        f"/subscriptions/{subscription_id}/resourceGroups/{resource_group_name}"
        f"/providers/{provider_path}"
    )


def build_resource_groups(
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    tags = common_tags(global_cfg)
    result: dict[str, Any] = {}
    for rg_id, row in rgs.items():
        require(row, "name", "location")
        result[rg_id] = {
            "name": str(row["name"]),
            "location": str(row["location"]),
            "tags": {**tags, **common_tags(global_cfg, row)},
        }
    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "location": global_cfg["default_location"],
        "common_tags": tags,
        "resource_groups": result,
    }


def build_hub(
    hub_id: str,
    hubs: dict[str, dict[str, Any]],
    hub_subnets: list[dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if hub_id not in hubs:
        raise DesignError(f"Unknown hub_id: {hub_id}")
    hub = hubs[hub_id]
    require(hub, "resource_group_id", "vnet_name", "location", "address_space")
    rg_id = str(hub["resource_group_id"])
    if rg_id not in rgs:
        raise DesignError(f"Hub {hub_id}: unknown resource_group_id={rg_id}")
    address_space = str(hub["address_space"])
    ensure_cidr(address_space, f"hub {hub_id}")
    subnets: dict[str, Any] = {}
    inbound_key = "dns_in"
    outbound_key = "dns_out"
    for row in hub_subnets:
        if not enabled(row.get("enabled", "Y")) or str(row.get("hub_id", "")) != hub_id:
            continue
        require(row, "subnet_id", "name", "address_prefix")
        subnet_id = str(row["subnet_id"])
        prefix = str(row["address_prefix"])
        ensure_cidr(prefix, f"hub subnet {subnet_id}")
        if not ip_network(prefix, strict=False).subnet_of(ip_network(address_space, strict=False)):
            raise DesignError(f"Hub subnet {subnet_id} is outside {address_space}")
        role = str(row.get("dns_endpoint_role", "")).lower()
        if role == "inbound":
            inbound_key = subnet_id
        elif role == "outbound":
            outbound_key = subnet_id
        subnets[subnet_id] = {
            "name": str(row["name"]),
            "address_prefixes": [prefix],
            "create_nsg": enabled(row.get("create_nsg")),
            "delegate_dns_resolver": enabled(row.get("delegate_dns_resolver")),
        }
    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "location": str(hub["location"]),
        "resource_group_name": str(rgs[rg_id]["name"]),
        "hub_name": str(hub["vnet_name"]),
        "address_space": [address_space],
        "dns_servers": split_list(hub.get("dns_servers")),
        "enable_private_dns_resolver": enabled(hub.get("private_dns_resolver")),
        "dns_inbound_subnet_key": inbound_key,
        "dns_outbound_subnet_key": outbound_key,
        "subnets": subnets,
        "common_tags": common_tags(global_cfg, hub),
    }


def build_spoke(
    workload_id: str,
    workloads: dict[str, dict[str, Any]],
    workload_subnets: list[dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if workload_id not in workloads:
        raise DesignError(f"Unknown workload_id: {workload_id}")
    workload = workloads[workload_id]
    require(workload, "resource_group_id", "vnet_name", "address_space", "location")
    rg_id = str(workload["resource_group_id"])
    if rg_id not in rgs:
        raise DesignError(f"Workload {workload_id}: unknown resource_group_id={rg_id}")
    address_space = str(workload["address_space"])
    ensure_cidr(address_space, f"workload {workload_id}")
    subnets: dict[str, Any] = {}
    for row in workload_subnets:
        if not enabled(row.get("enabled", "Y")) or str(row.get("workload_id", "")) != workload_id:
            continue
        require(row, "subnet_id", "name", "address_prefix")
        subnet_id = str(row["subnet_id"])
        prefix = str(row["address_prefix"])
        ensure_cidr(prefix, f"workload subnet {subnet_id}")
        if not ip_network(prefix, strict=False).subnet_of(ip_network(address_space, strict=False)):
            raise DesignError(f"Workload subnet {subnet_id} is outside {address_space}")
        subnets[subnet_id] = {
            "name": str(row["name"]),
            "address_prefixes": [prefix],
            "create_nsg": enabled(row.get("create_nsg")),
            "associate_route_table": enabled(row.get("associate_route_table")),
            "private_endpoint_network_policies": str(
                row.get("private_endpoint_network_policies", "Enabled")
            ),
        }
    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "location": str(workload["location"]),
        "resource_group_name": str(rgs[rg_id]["name"]),
        "spoke_name": str(workload["vnet_name"]),
        "address_space": [address_space],
        "dns_servers": [global_cfg["hub_dns_inbound_ip"]],
        "hub_resource_group_name": global_cfg["hub_resource_group_name"],
        "hub_vnet_name": global_cfg["hub_vnet_name"],
        "firewall_private_ip": global_cfg["firewall_private_ip"],
        "subnets": subnets,
        "common_tags": common_tags(global_cfg, workload),
    }


def build_vm(
    workload_id: str,
    workloads: dict[str, dict[str, Any]],
    workload_subnets: dict[str, dict[str, Any]],
    vm_rows: dict[str, dict[str, Any]],
    disk_rows: list[dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if workload_id not in workloads:
        raise DesignError(f"Unknown workload_id: {workload_id}")
    workload = workloads[workload_id]
    rg_id = str(workload["resource_group_id"])
    vm_disks: dict[str, dict[str, Any]] = defaultdict(dict)
    lun_seen: dict[str, set[int]] = defaultdict(set)
    for disk in disk_rows:
        if not enabled(disk.get("enabled", "Y")):
            continue
        require(disk, "disk_id", "vm_id", "disk_name", "size_gb", "lun", "storage_type", "caching")
        vm_id = str(disk["vm_id"])
        lun = int(disk["lun"])
        if lun in lun_seen[vm_id]:
            raise DesignError(f"VM {vm_id}: duplicate data disk LUN {lun}")
        lun_seen[vm_id].add(lun)
        vm_disks[vm_id][str(disk["disk_name"])] = {
            "size_gb": int(disk["size_gb"]),
            "storage_account_type": str(disk["storage_type"]),
            "lun": lun,
            "caching": str(disk["caching"]),
        }

    buckets: dict[str, dict[str, Any]] = {
        "web_vms": {},
        "was_vms": {},
        "db_vms": {},
        "agent_vms": {},
    }
    selected = [
        row for row in vm_rows.values()
        if str(row.get("workload_id", "")) == workload_id
    ]
    if not selected:
        raise DesignError(f"No enabled VMs for workload {workload_id}")
    first = selected[0]
    publisher, offer, sku = image_for(str(first.get("os", "RHEL9")))
    private_ips: set[str] = set()
    for row in selected:
        require(
            row, "vm_id", "role", "vm_name", "subnet_id", "private_ip",
            "vm_size", "os_disk_gb", "admin_username", "itsm_ticket",
        )
        vm_id = str(row["vm_id"])
        subnet_id = str(row["subnet_id"])
        if subnet_id not in workload_subnets:
            raise DesignError(f"VM {vm_id}: unknown subnet_id={subnet_id}")
        subnet = workload_subnets[subnet_id]
        if str(subnet.get("workload_id", "")) != workload_id:
            raise DesignError(f"VM {vm_id}: subnet {subnet_id} belongs to another workload")
        ip_value = str(row["private_ip"])
        if ip_value in private_ips:
            raise DesignError(f"Workload {workload_id}: duplicate private IP {ip_value}")
        private_ips.add(ip_value)
        ensure_ip_in_subnet(ip_value, str(subnet["address_prefix"]), f"VM {vm_id}")
        role = str(row["role"]).strip().lower()
        bucket = (
            "web_vms" if role == "web"
            else "was_vms" if role in {"was", "app"}
            else "db_vms" if role == "db"
            else "agent_vms"
        )
        buckets[bucket][vm_id] = {
            "name": str(row["vm_name"]),
            "subnet_name": str(subnet["name"]),
            "private_ip_address": ip_value,
            "vm_size": str(row["vm_size"]),
            "os_disk_size_gb": int(row["os_disk_gb"]),
            "role": role,
            "itsm_ticket": str(row["itsm_ticket"]),
            "data_disks": vm_disks.get(vm_id, {}),
        }
    return {
        "resource_group_name": str(rgs[rg_id]["name"]),
        "location": str(workload["location"]),
        "vnet_resource_group_name": str(rgs[rg_id]["name"]),
        "vnet_name": str(workload["vnet_name"]),
        "admin_username": str(first["admin_username"]),
        "ssh_public_key": global_cfg["ssh_public_key"],
        "image_publisher": publisher,
        "image_offer": offer,
        "image_sku": sku,
        "image_version": "latest",
        "os_disk_storage_type": "Premium_LRS",
        "common_tags": common_tags(global_cfg, workload),
        **buckets,
    }


def vnet_id_for_link(
    row: dict[str, Any],
    hubs: dict[str, dict[str, Any]],
    workloads: dict[str, dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> str:
    scope = str(row.get("vnet_scope", "")).strip().lower()
    key = str(row.get("vnet_key", "")).strip()
    if scope == "hub":
        if key not in hubs:
            raise DesignError(f"Private DNS link references unknown hub: {key}")
        item = hubs[key]
    elif scope == "workload":
        if key not in workloads:
            raise DesignError(f"Private DNS link references unknown workload: {key}")
        item = workloads[key]
    else:
        raise DesignError(f"Private DNS link vnet_scope must be hub or workload: {row}")

    rg_id = str(item.get("resource_group_id", ""))
    if rg_id not in rgs:
        raise DesignError(f"Private DNS link references unknown resource group: {rg_id}")
    require(item, "vnet_name")
    return resource_id(
        global_cfg["subscription_id"],
        str(rgs[rg_id]["name"]),
        f"Microsoft.Network/virtualNetworks/{item['vnet_name']}",
    )


def build_private_dns_zones(
    resource_group_id: str,
    dns_zones: dict[str, dict[str, Any]],
    dns_links: list[dict[str, Any]],
    hubs: dict[str, dict[str, Any]],
    workloads: dict[str, dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if resource_group_id not in rgs:
        raise DesignError(f"Private DNS generation references unknown resource group: {resource_group_id}")

    selected = {
        zone_id: row
        for zone_id, row in dns_zones.items()
        if str(row.get("resource_group_id", "")) == resource_group_id
    }
    if not selected:
        raise DesignError(
            f"No enabled Private DNS zones for resource_group_id={resource_group_id}"
        )

    zone_names: set[str] = set()
    for zone_id, row in selected.items():
        require(row, "zone_name")
        zone_name = str(row["zone_name"])
        if zone_name in zone_names:
            raise DesignError(f"Duplicate Private DNS zone name: {zone_name}")
        zone_names.add(zone_name)

    links: dict[str, Any] = {}
    for row in dns_links:
        if not enabled(row.get("enabled", "Y")):
            continue
        zone_id = str(row.get("dns_zone_id", "")).strip()
        if zone_id not in selected:
            continue
        require(row, "link_id", "link_name", "vnet_scope", "vnet_key")
        link_id = str(row["link_id"])
        if link_id in links:
            raise DesignError(f"Duplicate Private DNS link_id={link_id}")
        links[link_id] = {
            "zone_name": str(selected[zone_id]["zone_name"]),
            "name": str(row["link_name"]),
            "virtual_network_id": vnet_id_for_link(
                row, hubs, workloads, rgs, global_cfg
            ),
            "registration_enabled": enabled(row.get("registration_enabled")),
        }

    first = next(iter(selected.values()))
    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "resource_group_name": str(rgs[resource_group_id]["name"]),
        "zones": sorted(zone_names),
        "virtual_network_links": links,
        "common_tags": common_tags(global_cfg, first),
    }


def private_dns_zone_resource_id(
    zone_id_or_resource_id: str,
    dns_zones: dict[str, dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> str:
    if zone_id_or_resource_id.startswith("/subscriptions/"):
        return zone_id_or_resource_id
    if zone_id_or_resource_id not in dns_zones:
        raise DesignError(
            f"Private Endpoint references unknown Private DNS zone: {zone_id_or_resource_id}"
        )
    zone = dns_zones[zone_id_or_resource_id]
    rg_id = str(zone.get("resource_group_id", ""))
    if rg_id not in rgs:
        raise DesignError(f"Private DNS zone references unknown resource group: {rg_id}")
    require(zone, "zone_name")
    return resource_id(
        global_cfg["subscription_id"],
        str(rgs[rg_id]["name"]),
        f"Microsoft.Network/privateDnsZones/{zone['zone_name']}",
    )


def build_private_endpoint(
    endpoint_id: str,
    endpoints: dict[str, dict[str, Any]],
    dns_zones: dict[str, dict[str, Any]],
    workloads: dict[str, dict[str, Any]],
    workload_subnets: dict[str, dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if endpoint_id not in endpoints:
        raise DesignError(f"Unknown private_endpoint_id: {endpoint_id}")
    endpoint = endpoints[endpoint_id]
    require(
        endpoint,
        "resource_group_id", "workload_id", "subnet_id", "pe_name",
        "private_connection_resource_id", "subresource_names",
    )

    rg_id = str(endpoint["resource_group_id"])
    workload_id = str(endpoint["workload_id"])
    subnet_id = str(endpoint["subnet_id"])
    if rg_id not in rgs:
        raise DesignError(f"Private Endpoint references unknown resource group: {rg_id}")
    if workload_id not in workloads:
        raise DesignError(f"Private Endpoint references unknown workload: {workload_id}")
    if subnet_id not in workload_subnets:
        raise DesignError(f"Private Endpoint references unknown subnet: {subnet_id}")

    workload = workloads[workload_id]
    subnet = workload_subnets[subnet_id]
    if str(subnet.get("workload_id", "")) != workload_id:
        raise DesignError(
            f"Private Endpoint subnet {subnet_id} does not belong to {workload_id}"
        )
    workload_rg_id = str(workload.get("resource_group_id", ""))
    if workload_rg_id not in rgs:
        raise DesignError(
            f"Private Endpoint workload references unknown resource group: {workload_rg_id}"
        )

    subnet_resource_id = resource_id(
        global_cfg["subscription_id"],
        str(rgs[workload_rg_id]["name"]),
        (
            f"Microsoft.Network/virtualNetworks/{workload['vnet_name']}"
            f"/subnets/{subnet['name']}"
        ),
    )
    dns_zone_ids = [
        private_dns_zone_resource_id(item, dns_zones, rgs, global_cfg)
        for item in split_list(endpoint.get("private_dns_zone_ids"))
    ]
    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "location": str(workload.get("location", global_cfg["default_location"])),
        "resource_group_name": str(rgs[rg_id]["name"]),
        "pe_name": str(endpoint["pe_name"]),
        "subnet_id": subnet_resource_id,
        "private_connection_resource_id": str(
            endpoint["private_connection_resource_id"]
        ),
        "subresource_names": split_list(endpoint["subresource_names"]),
        "private_dns_zone_ids": dns_zone_ids,
        "common_tags": common_tags(global_cfg, endpoint),
    }


def build_firewall_rule(
    rule_group_id: str,
    rule_groups: dict[str, dict[str, Any]],
    network_rule_rows: list[dict[str, Any]],
    application_rule_rows: list[dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if rule_group_id not in rule_groups:
        raise DesignError(f"Unknown firewall rule_group_id: {rule_group_id}")
    group = rule_groups[rule_group_id]
    require(
        group, "firewall_policy_id", "rule_collection_group_name", "priority"
    )

    network_collections: dict[str, Any] = {}
    for row in network_rule_rows:
        if not enabled(row.get("enabled", "Y")):
            continue
        if str(row.get("rule_group_id", "")) != rule_group_id:
            continue
        require(
            row, "rule_id", "collection_id", "collection_name",
            "collection_priority", "action", "name", "protocols",
            "source_addresses", "destination_addresses", "destination_ports",
        )
        collection_id = str(row["collection_id"])
        collection = network_collections.setdefault(
            collection_id,
            {
                "name": str(row["collection_name"]),
                "priority": int(row["collection_priority"]),
                "action": str(row["action"]),
                "rules": {},
            },
        )
        if (
            collection["name"] != str(row["collection_name"])
            or collection["priority"] != int(row["collection_priority"])
            or collection["action"] != str(row["action"])
        ):
            raise DesignError(
                f"Inconsistent network collection metadata: {collection_id}"
            )
        rule_id = str(row["rule_id"])
        if rule_id in collection["rules"]:
            raise DesignError(f"Duplicate firewall network rule_id={rule_id}")
        collection["rules"][rule_id] = {
            "name": str(row["name"]),
            "protocols": split_list(row["protocols"]),
            "source_addresses": split_list(row["source_addresses"]),
            "destination_addresses": split_list(row["destination_addresses"]),
            "destination_ports": split_list(row["destination_ports"]),
        }

    application_collections: dict[str, Any] = {}
    for row in application_rule_rows:
        if not enabled(row.get("enabled", "Y")):
            continue
        if str(row.get("rule_group_id", "")) != rule_group_id:
            continue
        require(
            row, "rule_id", "collection_id", "collection_name",
            "collection_priority", "action", "name", "source_addresses",
            "protocol_type", "protocol_port", "destination_fqdns",
        )
        collection_id = str(row["collection_id"])
        collection = application_collections.setdefault(
            collection_id,
            {
                "name": str(row["collection_name"]),
                "priority": int(row["collection_priority"]),
                "action": str(row["action"]),
                "rules": {},
            },
        )
        if (
            collection["name"] != str(row["collection_name"])
            or collection["priority"] != int(row["collection_priority"])
            or collection["action"] != str(row["action"])
        ):
            raise DesignError(
                f"Inconsistent application collection metadata: {collection_id}"
            )
        rule_id = str(row["rule_id"])
        if rule_id in collection["rules"]:
            raise DesignError(f"Duplicate firewall application rule_id={rule_id}")
        collection["rules"][rule_id] = {
            "name": str(row["name"]),
            "source_addresses": split_list(row["source_addresses"]),
            "protocol_type": str(row["protocol_type"]),
            "protocol_port": int(row["protocol_port"]),
            "destination_fqdns": split_list(row["destination_fqdns"]),
        }

    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "location": global_cfg["default_location"],
        "common_tags": common_tags(global_cfg, group),
        "firewall_policy_id": str(group["firewall_policy_id"]),
        "rule_collection_group_name": str(
            group["rule_collection_group_name"]
        ),
        "priority": int(group["priority"]),
        "network_rule_collections": network_collections,
        "application_rule_collections": application_collections,
    }


def build_dns_records(
    dns_zone_id: str,
    dns_zones: dict[str, dict[str, Any]],
    dns_record_rows: list[dict[str, Any]],
    rgs: dict[str, dict[str, Any]],
    global_cfg: dict[str, str],
) -> dict[str, Any]:
    if dns_zone_id not in dns_zones:
        raise DesignError(f"Unknown dns_zone_id for records: {dns_zone_id}")
    zone = dns_zones[dns_zone_id]
    rg_id = str(zone.get("resource_group_id", ""))
    if rg_id not in rgs:
        raise DesignError(f"DNS record zone references unknown resource group: {rg_id}")
    require(zone, "zone_name")

    records: dict[str, Any] = {}
    first_row: dict[str, Any] | None = None
    for row in dns_record_rows:
        if not enabled(row.get("enabled", "Y")):
            continue
        if str(row.get("dns_zone_id", "")) != dns_zone_id:
            continue
        require(row, "record_id", "name", "records")
        record_id = str(row["record_id"])
        if record_id in records:
            raise DesignError(f"Duplicate DNS record_id={record_id}")
        records[record_id] = {
            "name": str(row["name"]),
            "ttl": int(row.get("ttl") or 300),
            "records": split_list(row["records"]),
        }
        first_row = first_row or row
    if not records:
        raise DesignError(f"No enabled DNS records for dns_zone_id={dns_zone_id}")

    return {
        "tenant_id": global_cfg["tenant_id"],
        "subscription_id": global_cfg["subscription_id"],
        "location": global_cfg["default_location"],
        "resource_group_name": str(rgs[rg_id]["name"]),
        "zone_name": str(zone["zone_name"]),
        "a_records": records,
        "common_tags": common_tags(global_cfg, first_row or zone),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--excel", default="design/azure_landingzone_design.xlsx")
    parser.add_argument("--out", default="generated_tfvars")
    parser.add_argument("--tenant-id", default=os.getenv("TENANT_ID", ""))
    parser.add_argument("--subscription-id", default=os.getenv("SUBSCRIPTION_ID", ""))
    parser.add_argument("--ssh-public-key", default=os.getenv("SSH_PUBLIC_KEY", ""))
    parser.add_argument(
        "--clean",
        action="store_true",
        help="Remove the output directory before generation",
    )
    args = parser.parse_args()

    workbook = load_workbook(args.excel, data_only=True)
    required_sheets = [
        "01_Global",
        "02_ResourceGroups",
        "03_HubNetwork",
        "04_HubSubnets",
        "05_Workloads",
        "06_WorkloadSubnets",
        "07_VMs",
        "08_VMDisks",
        "09_PrivateDNSZones",
        "10_PrivateDNSLinks",
        "11_PrivateEndpoints",
        "12_FirewallRuleGroups",
        "13_FirewallNetworkRules",
        "14_FirewallApplicationRules",
        "15_DNSRecords",
        "99_GenerationMap",
    ]
    missing = [name for name in required_sheets if name not in workbook.sheetnames]
    if missing:
        raise SystemExit(f"Missing Excel sheets: {missing}")

    global_cfg = kv_sheet(workbook["01_Global"])
    context = {
        "TENANT_ID": args.tenant_id,
        "SUBSCRIPTION_ID": args.subscription_id,
        "SSH_PUBLIC_KEY": args.ssh_public_key,
    }
    for key in ["tenant_id", "subscription_id", "ssh_public_key"]:
        global_cfg[key] = str(
            resolve_placeholders(global_cfg.get(key, ""), context)
        )
    required_globals = [
        "tenant_id",
        "subscription_id",
        "default_location",
        "project",
        "managed_by",
        "hub_resource_group_name",
        "hub_vnet_name",
        "hub_dns_inbound_ip",
        "firewall_private_ip",
        "ssh_public_key",
    ]
    empty_globals = [key for key in required_globals if not global_cfg.get(key)]
    if empty_globals:
        raise SystemExit(
            f"Missing global values or environment variables: {empty_globals}"
        )

    rgs = unique(
        rows(workbook["02_ResourceGroups"]),
        "resource_group_id",
        "02_ResourceGroups",
    )
    hubs = unique(rows(workbook["03_HubNetwork"]), "hub_id", "03_HubNetwork")
    workloads = unique(
        rows(workbook["05_Workloads"]),
        "workload_id",
        "05_Workloads",
    )
    workload_subnets = unique(
        rows(workbook["06_WorkloadSubnets"]),
        "subnet_id",
        "06_WorkloadSubnets",
    )
    vms = unique(rows(workbook["07_VMs"]), "vm_id", "07_VMs")
    dns_zones = unique(
        rows(workbook["09_PrivateDNSZones"]),
        "dns_zone_id",
        "09_PrivateDNSZones",
    )
    private_endpoints = unique(
        rows(workbook["11_PrivateEndpoints"]),
        "private_endpoint_id",
        "11_PrivateEndpoints",
    )
    firewall_rule_groups = unique(
        rows(workbook["12_FirewallRuleGroups"]),
        "rule_group_id",
        "12_FirewallRuleGroups",
    )
    generation_map = unique(
        rows(workbook["99_GenerationMap"]),
        "generation_id",
        "99_GenerationMap",
    )

    hub_subnets = rows(workbook["04_HubSubnets"])
    disks = rows(workbook["08_VMDisks"])
    dns_links = rows(workbook["10_PrivateDNSLinks"])
    firewall_network_rules = rows(workbook["13_FirewallNetworkRules"])
    firewall_application_rules = rows(
        workbook["14_FirewallApplicationRules"]
    )
    dns_records = rows(workbook["15_DNSRecords"])

    output = Path(args.out)
    if args.clean and output.exists():
        shutil.rmtree(output)

    for generation_id, item in generation_map.items():
        require(
            item,
            "module_type",
            "source_key",
            "output_root",
            "output_filename",
        )
        module_type = str(item["module_type"]).strip()
        source_key = str(item["source_key"]).strip()
        path = output / str(item["output_root"]) / str(item["output_filename"])

        if module_type == "resource_groups":
            data = build_resource_groups(rgs, global_cfg)
        elif module_type == "hub_network":
            data = build_hub(
                source_key, hubs, hub_subnets, rgs, global_cfg
            )
        elif module_type == "workload_spoke":
            data = build_spoke(
                source_key,
                workloads,
                list(workload_subnets.values()),
                rgs,
                global_cfg,
            )
        elif module_type == "vm":
            data = build_vm(
                source_key,
                workloads,
                workload_subnets,
                vms,
                disks,
                rgs,
                global_cfg,
            )
        elif module_type == "private_dns_zones":
            data = build_private_dns_zones(
                source_key,
                dns_zones,
                dns_links,
                hubs,
                workloads,
                rgs,
                global_cfg,
            )
        elif module_type == "private_endpoint":
            data = build_private_endpoint(
                source_key,
                private_endpoints,
                dns_zones,
                workloads,
                workload_subnets,
                rgs,
                global_cfg,
            )
        elif module_type == "firewall_rule":
            data = build_firewall_rule(
                source_key,
                firewall_rule_groups,
                firewall_network_rules,
                firewall_application_rules,
                global_cfg,
            )
        elif module_type == "dns_records":
            data = build_dns_records(
                source_key,
                dns_zones,
                dns_records,
                rgs,
                global_cfg,
            )
        else:
            raise DesignError(
                f"{generation_id}: unsupported module_type={module_type}"
            )
        write_json(path, data, context)

    print("Excel design validation and tfvars generation completed.")


if __name__ == "__main__":
    try:
        main()
    except DesignError as exc:
        raise SystemExit(f"DESIGN ERROR: {exc}") from exc
