#!/usr/bin/env python3
"""Prepare the Excel design file used by the Azure DevOps pipeline.

When --secure-file is supplied, the protected Excel file is copied to the
pipeline working location. Otherwise, a default workbook is generated with
create_design_excel.py. The resulting workbook is validated before use.
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_SHEETS = {
    "00_Control",
    "01_Global",
    "02_ResourceGroups",
    "03_HubNetwork",
    "04_HubSubnets",
    "05_Workloads",
    "06_WorkloadSubnets",
    "07_VMs",
    "08_VMDisks",
    "09_PrivateDNSZones",
    "10_PrivateDNSLinks",
    "11_PrivateEndpoints",
    "12_FirewallRuleGroups",
    "13_FirewallNetworkRules",
    "14_FirewallApplicationRules",
    "15_DNSRecords",
    "99_GenerationMap",
}


def validate_workbook(path: Path) -> None:
    """Fail when the workbook is missing or does not match the design schema."""
    if not path.is_file():
        raise SystemExit(f"Excel design file was not created: {path}")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        available = set(workbook.sheetnames)
        missing = sorted(REQUIRED_SHEETS - available)
        if missing:
            raise SystemExit(
                "Excel design workbook is missing required sheets: "
                + ", ".join(missing)
            )
    finally:
        workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--out",
        required=True,
        help="Destination path for the prepared Excel workbook",
    )
    parser.add_argument(
        "--secure-file",
        help="Optional Azure DevOps Secure File path to copy instead of generating",
    )
    args = parser.parse_args()

    output = Path(args.out).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    if args.secure_file:
        source = Path(args.secure_file).expanduser().resolve()
        if not source.is_file():
            raise SystemExit(f"Secure Excel design file does not exist: {source}")
        shutil.copy2(source, output)
        print(f"copied secure Excel design: {source} -> {output}")
    else:
        generator = Path(__file__).with_name("create_design_excel.py").resolve()
        if not generator.is_file():
            raise SystemExit(f"Excel generator was not found: {generator}")

        subprocess.run(
            [
                sys.executable,
                str(generator),
                "--out",
                str(output),
                "--force",
            ],
            check=True,
        )
        print(f"generated default Excel design: {output}")

    validate_workbook(output)
    print(f"validated Excel design: {output}")


if __name__ == "__main__":
    main()
